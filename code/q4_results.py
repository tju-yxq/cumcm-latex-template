# -*- coding: utf-8 -*-
"""问题4 结果输出：result4-2.xlsx（问题二格式）与 result4-3.xlsx（问题三格式）、图表。"""
import sys, json
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent))
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from common import load_all, key_numbers, log, RESULTS, FIGURES, N, DT, DAYS
from xlsx_writer import (write_plan_sheet, write_battery_sheet, write_emergency_sheet,
                         spec_date_detail, soc_boundaries)

A1, LOAD, PV, PRICE4, FC3 = load_all()

# ---- result4-2.xlsx（问题二格式：计划/充放电/紧急）----
s2 = np.load(RESULTS / "q4_2_sim.npz")
tpl2 = r"C:\Users\26517\Documents\CUMCM\C题\附件\附件5\result4-2.xlsx"
wb = load_workbook(tpl2)
write_plan_sheet(wb["计划购电量"], s2["g_plan"], s2["cost_plan"])
soc_0, soc_24 = soc_boundaries(s2["c"], s2["d"], s2["soc_start"])
write_battery_sheet(wb["充放电量"], s2["c"], s2["d"], soc_0, soc_24)
n2 = write_emergency_sheet(wb["紧急购电量"], s2["emg"])
wb.save(RESULTS / "result4-2.xlsx")
log(f"result4-2.xlsx 已写出（{n2} 个紧急购电日）")

# ---- result4-3.xlsx（问题三格式：计划/调整/充放电/紧急）----
s3 = np.load(RESULTS / "q4_3_sim.npz")
tpl3 = r"C:\Users\26517\Documents\CUMCM\C题\附件\附件5\result4-3.xlsx"
wb = load_workbook(tpl3)
plan_cost = np.array([float(PRICE4[d] @ s3["g_plan"][d]) for d in range(DAYS)])
write_plan_sheet(wb["计划购电量"], s3["g_plan"], plan_cost)
write_plan_sheet(wb["调整购电量"], s3["a_final"], s3["cost_settle"])
soc_0b, soc_24b = soc_boundaries(s3["c"], s3["d"], s3["soc_start"])
write_battery_sheet(wb["充放电量"], s3["c"], s3["d"], soc_0b, soc_24b)
n3 = write_emergency_sheet(wb["紧急购电量"], s3["emg"])
wb.save(RESULTS / "result4-3.xlsx")
log(f"result4-3.xlsx 已写出（{n3} 个紧急购电日）")

# ---- 四指定日明细（P4-3）----
def cost_by_day(D):
    return dict(total_g=round(float(s3["a_final"][D].sum()), 2),
                cost_settle=round(float(s3["cost_settle"][D]), 2),
                cost_emg=round(float(s3["cost_emg"][D]), 2),
                cost_total=round(float(s3["cost_settle"][D] + s3["cost_emg"][D]), 2),
                adj_qty=round(float(np.abs(s3["a_final"][D] - s3["g_plan"][D]).sum()), 2))

detail = spec_date_detail(s3["g_plan"], s3["a_final"], s3["c"], s3["d"], s3["emg"],
                          s3["soc_start"], cost_by_day)
for s, v in detail.items():
    print(f"P4-3 {s}: 最终购电 {v['total_g']:.0f} kWh, 结算 {v['cost_settle']:.0f}, "
          f"紧急 {v['cost_emg']:.0f}, 调整量 {v['adj_qty']:.0f}, 紧急时段 {len(v['t3_emergency'])}")
key_numbers(q4_3_spec_dates=detail)

# ---- 图：固定电价 vs 波动电价的策略费用对比 ----
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
plt.rcParams["font.sans-serif"] = ["SimHei", "Microsoft YaHei"]
plt.rcParams["axes.unicode_minus"] = False
kn = json.loads((RESULTS / "key_numbers.json").read_text(encoding="utf-8"))

labels = ["无储能基准", "问题二模型\n(P4-2)", "问题三模型\n不调整", "问题三模型\n全调整(P4-3)", "理想先知"]
fixed = [kn["q2_cost_baseline"], kn["q2_cost_main"], kn["q3_cost_noAdj"],
         kn["q3_cost_main"], kn["q2_cost_oracle"]]
fluct = [kn["q4_cost_baseline"], kn["q4_2_cost_main"], kn["q4_3_cost_noAdj"],
         kn["q4_3_cost_main"], kn["q4_2_cost_oracle"]]
x = np.arange(5); w = 0.36
fig, ax = plt.subplots(figsize=(9.2, 4.6))
ax.bar(x - w / 2, np.array(fixed) / 1e4, w, label="固定电价（附件1）", color="#5d6d7e")
ax.bar(x + w / 2, np.array(fluct) / 1e4, w, label="波动电价（附件4）", color="#c0392b")
for xi, v in zip(x - w / 2, fixed):
    ax.text(xi, v / 1e4 + 1.0, f"{v/1e4:.0f}", ha="center", fontsize=8.5)
for xi, v in zip(x + w / 2, fluct):
    ax.text(xi, v / 1e4 + 1.0, f"{v/1e4:.0f}", ha="center", fontsize=8.5)
ax.set_xticks(x); ax.set_xticklabels(labels, fontsize=9)
ax.set_ylabel("回测期总费用（万元）")
ax.set_title("固定电价与波动电价下各策略总费用对比（2025.2.1–12.31）")
ax.legend()
fig.tight_layout(); fig.savefig(FIGURES / "fig_q4_cmp.png", dpi=300); plt.close(fig)
log("图已输出：fig_q4_cmp")
