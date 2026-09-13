# -*- coding: utf-8 -*-
"""问题2 结果输出：result2.xlsx（按附件5模板）、四指定日明细、论文图表。"""
import sys, io, json
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent))
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from common import load_all, key_numbers, log, RESULTS, FIGURES, N, DT, ETA, SOC_INIT, DAYS
from q2_simulate import clock_day_view, emergency_intervals, PRICE_A1, REP

A1, LOAD, PV, PRICE4, FC3 = load_all()
sim = np.load(RESULTS / "q2_main_sim.npz")
g_plan, C, D_, EM = sim["g_plan"], sim["c"], sim["d"], sim["emg"]
soc_start = sim["soc_start"]
cost_plan, cost_emg = sim["cost_plan"], sim["cost_emg"]

DATES = pd.date_range("2025-01-01", periods=DAYS)
REP_DATES = DATES[31:365]                      # 2025-02-01 .. 12-31（334天）

# ---- 钟表日视图（执行量）----
c_ck = clock_day_view(C); d_ck = clock_day_view(D_); emg_ck = clock_day_view(EM)
# 钟表日 SOC 边界：SOC(D 0:00) = soc_start[D] − Δ(计划日D−1末槽)
soc_0 = np.zeros(DAYS); soc_24 = np.zeros(DAYS)
soc_0[0] = SOC_INIT
for D in range(1, DAYS):
    soc_0[D] = soc_start[D] - (ETA * C[D - 1, 143] - D_[D - 1, 143] / ETA)
soc_24 = np.roll(soc_0, -1)
# 12月31日 24:00 = 计划日364第143槽（k=142）之后
soc_24[364] = soc_start[364] + float(np.sum(ETA * C[364, :143] - D_[364, :143] / ETA))
assert abs(soc_24[363] - soc_0[364]) < 1e-6

BLOCKS = [(0, 24, "0:00-4:00"), (24, 48, "4:00-8:00"), (48, 72, "8:00-12:00"),
          (72, 96, "12:00-16:00"), (96, 120, "16:00-20:00"), (120, 144, "20:00-24:00")]

# ---- 写 result2.xlsx ----
tpl = r"C:\Users\26517\Documents\CUMCM\C题\附件\附件5\result2.xlsx"
wb = load_workbook(tpl)
ws = wb["计划购电量"]
hdr = [ws.cell(row=1, column=c).value for c in range(1, 148)]
assert hdr[1] == "0:10-0:20" and hdr[144] == "0:00-0:10+1", f"模板列标签异常: {hdr[1]}, {hdr[144]}"
ws.delete_rows(2, ws.max_row)
for i, D in enumerate(range(31, 365)):
    r = i + 2
    ws.cell(row=r, column=1, value=REP_DATES[i].to_pydatetime())
    for k in range(144):
        ws.cell(row=r, column=2 + k, value=round(float(g_plan[D, k]), 4))
    ws.cell(row=r, column=146, value=round(float(g_plan[D].sum()), 2))
    ws.cell(row=r, column=147, value=round(float(cost_plan[D]), 2))

ws2 = wb["充放电量"]
ws2.delete_rows(2, ws2.max_row)
r = 2
for i, D in enumerate(range(31, 365)):
    for b, (a, bnd, lab) in enumerate(BLOCKS):
        ws2.cell(row=r, column=1, value=REP_DATES[i].to_pydatetime() if b == 0 else None)
        ws2.cell(row=r, column=2, value=lab)
        ws2.cell(row=r, column=3, value=round(float(c_ck[D, a:bnd].sum()), 4))
        ws2.cell(row=r, column=4, value=round(float(d_ck[D, a:bnd].sum()), 4))
        if b == 0:
            ws2.cell(row=r, column=5, value="0:00"); ws2.cell(row=r, column=6, value=round(float(soc_0[D]), 2))
        elif b == 1:
            ws2.cell(row=r, column=5, value="24:00"); ws2.cell(row=r, column=6, value=round(float(soc_24[D]), 2))
        r += 1

ws3 = wb["紧急购电量"]
ws3.delete_rows(2, ws3.max_row)
r = 2
emg_days = 0
for D in range(31, 365):
    ivs = emergency_intervals(emg_ck, D)
    if not ivs:
        continue
    emg_days += 1
    for j, (lab, q) in enumerate(ivs):
        ws3.cell(row=r, column=1, value=REP_DATES[D - 31].to_pydatetime() if j == 0 else None)
        ws3.cell(row=r, column=2, value=lab)
        ws3.cell(row=r, column=3, value=round(q, 2))
        r += 1
out = RESULTS / "result2.xlsx"
wb.save(out)
log(f"result2.xlsx 已写出：334天计划 + {334*6}行充放电 + {emg_days}个紧急购电日")

# ---- 四指定日明细（论文表3用）----
SPEC = ["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]
detail = {}
for s in SPEC:
    D = (pd.Timestamp(s) - pd.Timestamp("2025-01-01")).days
    t1_slots = [60, 72, 84, 96, 108, 120]     # 10:00,12:00,...,20:00（钟表槽）
    # 计划购电量按计划日（文件列序 label=(k+1)*10 分钟）取槽 10:00..20:00 → k=59,71,...
    t1 = {f"{(k+1)*10//60:02d}:{(k+1)*10%60:02d}": round(float(g_plan[D, k]), 2)
          for k in [59, 71, 83, 95, 107, 119]}
    t2 = {lab: [round(float(c_ck[D, a:bnd].sum()), 2), round(float(d_ck[D, a:bnd].sum()), 2)]
          for a, bnd, lab in BLOCKS}
    t3 = emergency_intervals(emg_ck, D)
    detail[s] = dict(
        day_index=D, total_g=round(float(g_plan[D].sum()), 2),
        cost_plan=round(float(cost_plan[D]), 2), cost_emg=round(float(cost_emg[D]), 2),
        cost_total=round(float(cost_plan[D] + cost_emg[D]), 2),
        soc_0=round(float(soc_0[D]), 2), soc_24=round(float(soc_24[D]), 2),
        t1_slots=t1, t2_blocks=t2, t3_emergency=t3)
    print(f"{s}: 计划 {detail[s]['total_g']:.0f} kWh / {detail[s]['cost_plan']:.0f} 元, "
          f"紧急 {detail[s]['cost_emg']:.0f} 元, SOC {detail[s]['soc_0']:.0f}->{detail[s]['soc_24']:.0f}, "
          f"紧急时段 {len(t3)} 段")

# ---- 图 ----
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
plt.rcParams["font.sans-serif"] = ["SimHei", "Microsoft YaHei"]
plt.rcParams["axes.unicode_minus"] = False

# 图A：策略对比（重算各策略数值从 key_numbers 读取）
kn = json.loads((RESULTS / "key_numbers.json").read_text(encoding="utf-8"))
strategies = ["无储能\n基准", "朴素预测\n+贪心", "分位数缓冲\n+贪心", "分位数缓冲\n+滚动(主)", "理想先知"]
totals = [kn["q2_cost_baseline"], kn["q2_cost_naive"], kn["q2_cost_greedy"],
          kn["q2_cost_main"], kn["q2_cost_oracle"]]
fig, ax = plt.subplots(figsize=(8.6, 4.6))
x = np.arange(len(strategies))
bars = ax.bar(x, [t / 1e4 for t in totals], 0.58,
              color=["#95a5a6", "#e67e22", "#3498db", "#c0392b", "#27ae60"])
for xi, t in zip(x, totals):
    ax.text(xi, t / 1e4 + 12, f"{t/1e4:.1f}", ha="center", fontsize=10)
ax.set_xticks(x); ax.set_xticklabels(strategies, fontsize=9.5)
ax.set_ylabel("回测期总购电费用（万元）")
ax.set_title("问题二：各策略全年购电费用对比（2025.2.1–12.31）")
ax.set_ylim(0, max(totals) / 1e4 * 1.15)
fig.tight_layout()
fig.savefig(FIGURES / "fig_q2_strategy_cmp.png", dpi=300); plt.close(fig)

# 图B：缓冲分位数敏感性
sens = kn["q2_sens_q"]
qs = sorted(float(k) for k in sens)
fig, ax = plt.subplots(figsize=(7.4, 4.2))
ax.plot([q * 100 for q in qs], [sens[str(q)] / 1e4 for q in qs], "o-", color="#2c3e50", lw=1.8)
ax.axvline(80, color="#c0392b", ls="--", lw=1, label="报童理论值 80%")
ax.set_xlabel("缓冲分位数 q（%）"); ax.set_ylabel("回测期总费用（万元）")
ax.set_title("缓冲分位数敏感性"); ax.legend()
fig.tight_layout(); fig.savefig(FIGURES / "fig_q2_sens_q.png", dpi=300); plt.close(fig)

# 图C：典型紧急日调度图（正式回测期内紧急费用最大的日期）
emg_cost_daily = cost_emg.copy(); emg_cost_daily[:31] = 0.0    # 排除1月过渡期
Dmax = int(np.argmax(emg_cost_daily))
lab_time = [(k + 1) * 10 / 60 for k in range(N)]
fig, axes = plt.subplots(2, 1, figsize=(9.5, 6.2), sharex=True,
                         gridspec_kw={"height_ratios": [3, 1.6]})
ax = axes[0]
ax.plot(lab_time, PRICE_A1, color="#888888", lw=1.2, label="电价（右轴）")
ax.set_ylabel("电价（元/kWh）", color="#888888"); ax.tick_params(axis="y", labelcolor="#888888")
ax2 = ax.twinx()
ax2.step(lab_time, g_plan[Dmax], where="post", color="#c0392b", lw=1.5, label="计划购电量")
ax2.bar(lab_time, C[Dmax], width=10/60, align="edge", color="#2471a3", alpha=0.55, label="实际充电")
ax2.bar(lab_time, -D_[Dmax], width=10/60, align="edge", color="#1e8449", alpha=0.55, label="实际放电")
emg_slots = EM[Dmax] > 1e-6
ax2.bar(np.array(lab_time)[emg_slots], EM[Dmax][emg_slots], width=10/60, align="edge",
        color="#8e44ad", alpha=0.9, label="紧急购电")
ax2.axhline(0, color="k", lw=0.6); ax2.set_ylabel("电量（kWh/10min）")
h1, l1 = ax.get_legend_handles_labels(); h2, l2 = ax2.get_legend_handles_labels()
ax.legend(h1 + h2, l1 + l2, loc="upper left", fontsize=8.5, ncol=5, framealpha=0.9)
ax.set_title(f"问题二主策略在紧急购电最严重日的调度（{DATES[Dmax]:%Y-%m-%d}）")
soc_tr = soc_start[Dmax] + np.cumsum(ETA * C[Dmax] - D_[Dmax] / ETA)
axes[1].plot(lab_time, soc_tr, color="#6c3483", lw=1.8)
axes[1].axhline(1200, color="r", ls="--", lw=0.8); axes[1].axhline(10800, color="r", ls="--", lw=0.8)
axes[1].set_ylabel("储电量（kWh）"); axes[1].set_xlabel("时刻（h）"); axes[1].set_ylim(0, 12000)
fig.tight_layout(); fig.savefig(FIGURES / "fig_q2_worst_day.png", dpi=300); plt.close(fig)
log(f"图已输出：fig_q2_strategy_cmp / fig_q2_sens_q / fig_q2_worst_day({DATES[Dmax]:%m-%d})")

key_numbers(q2_spec_dates=detail, q2_worst_day=DATES[Dmax].strftime("%Y-%m-%d"),
            q2_emg_days=emg_days)
log("四指定日明细与关键数值已保存")
